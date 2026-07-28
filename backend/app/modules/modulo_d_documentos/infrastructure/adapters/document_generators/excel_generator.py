import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.modulo_d_documentos.domain.ports.reporte_generator_port import ReporteGeneratorPort

_SLATE_900 = "0F172A"
_SLATE_700 = "334155"
_SLATE_500 = "64748B"
_SLATE_100 = "F1F5F9"
_SLATE_50 = "F8FAFC"
_WHITE = "FFFFFF"
_EMERALD = "047857"
_ROSE = "BE123C"
_ROSE_HEADER = "E11D48"
_INDIGO = "4338CA"
_INDIGO_HEADER = "4F46E5"
_AMBER = "B45309"

_MONEDA = '"S/" #,##0.00'
_ENTERO = "#,##0"
_PORCENTAJE = "0.0%"

_LADO = Side(style="thin", color="E2E8F0")
_LADO_TOTAL = Side(style="double", color="94A3B8")
_BORDE = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)
_BORDE_TOTAL = Border(left=_LADO, right=_LADO, top=_LADO_TOTAL, bottom=_LADO)


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


class ExcelGenerator(ReporteGeneratorPort):
    async def generar_excel(self, datos: dict, nombre: str) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.sheet_view.showGridLines = False

        if "egresos" in datos:
            self._render_egresos(ws, datos)
        elif "total_vendido" in datos:
            self._render_resumen(ws, datos)
        elif "productos" in datos:
            self._render_mas_vendidos(ws, datos)
        else:
            self._cabecera(ws, datos.get("titulo", "Reporte"), datos.get("periodo", ""), 3)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def _cabecera(self, ws: Worksheet, titulo: str, periodo: str, ncols: int) -> int:
        """Banda de título + subtítulo. Devuelve la primera fila libre."""
        ultima = get_column_letter(ncols)

        ws.merge_cells(f"A1:{ultima}1")
        titulo_cell = ws["A1"]
        titulo_cell.value = titulo
        titulo_cell.font = Font(bold=True, size=16, color=_WHITE)
        titulo_cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, ncols + 1):
            ws.cell(row=1, column=col).fill = _fill(_SLATE_900)
        ws.row_dimensions[1].height = 34

        ws.merge_cells(f"A2:{ultima}2")
        generado = datetime.now().strftime("%d/%m/%Y %H:%M")
        sub_cell = ws["A2"]
        sub_cell.value = f"Período:  {periodo}          •          Generado:  {generado}"
        sub_cell.font = Font(size=10, color=_WHITE)
        sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, ncols + 1):
            ws.cell(row=2, column=col).fill = _fill(_SLATE_700)
        ws.row_dimensions[2].height = 20

        return 4

    def _seccion(self, ws: Worksheet, fila: int, texto: str, ncols: int, color: str = _INDIGO_HEADER) -> int:
        ultima = get_column_letter(ncols)
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        cell = ws.cell(row=fila, column=1, value=texto)
        cell.font = Font(bold=True, size=12, color=_WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, ncols + 1):
            ws.cell(row=fila, column=col).fill = _fill(color)
        ws.row_dimensions[fila].height = 24
        return fila + 1

    def _kpi(
        self,
        ws: Worksheet,
        fila: int,
        etiqueta: str,
        valor,
        ncols: int,
        fmt: str | None = None,
        color_valor: str = _SLATE_900,
        indice: int = 0,
    ) -> int:
        """Fila etiqueta/valor con zebra; el valor se extiende hasta la última columna."""
        fondo = _SLATE_50 if indice % 2 else _WHITE
        for col in range(1, ncols + 1):
            celda = ws.cell(row=fila, column=col)
            celda.fill = _fill(fondo)
            celda.border = _BORDE

        etiqueta_cell = ws.cell(row=fila, column=1, value=etiqueta)
        etiqueta_cell.font = Font(bold=True, size=11, color=_SLATE_700)
        etiqueta_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        valor_cell = ws.cell(row=fila, column=2, value=valor)
        valor_cell.font = Font(bold=True, size=11, color=color_valor)
        valor_cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        if fmt:
            valor_cell.number_format = fmt

        if ncols > 2:
            ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=ncols)
        ws.row_dimensions[fila].height = 20
        return fila + 1

    def _encabezado_tabla(
        self,
        ws: Worksheet,
        fila: int,
        headers: list[str],
        aligns: list[str] | None = None,
        color: str = _INDIGO_HEADER,
    ) -> int:
        for col, texto in enumerate(headers, 1):
            celda = ws.cell(row=fila, column=col, value=texto)
            celda.font = Font(bold=True, size=11, color=_WHITE)
            celda.fill = _fill(color)
            celda.border = _BORDE
            celda.alignment = Alignment(
                horizontal=(aligns[col - 1] if aligns else "left"), vertical="center", indent=1
            )
        ws.row_dimensions[fila].height = 22
        return fila + 1

    def _fila_tabla(
        self,
        ws: Worksheet,
        fila: int,
        valores: list,
        fmts: list[str | None] | None = None,
        aligns: list[str] | None = None,
        indice: int = 0,
    ) -> int:
        fondo = _SLATE_50 if indice % 2 else _WHITE
        for col, valor in enumerate(valores, 1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.fill = _fill(fondo)
            celda.border = _BORDE
            celda.font = Font(size=11, color=_SLATE_700)
            celda.alignment = Alignment(
                horizontal=(aligns[col - 1] if aligns else "left"), vertical="center", indent=1
            )
            if fmts and fmts[col - 1]:
                celda.number_format = fmts[col - 1]
        return fila + 1

    def _fila_total(
        self,
        ws: Worksheet,
        fila: int,
        valores: list,
        fmts: list[str | None] | None = None,
        aligns: list[str] | None = None,
    ) -> int:
        for col, valor in enumerate(valores, 1):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.fill = _fill(_SLATE_100)
            celda.border = _BORDE_TOTAL
            celda.font = Font(bold=True, size=11, color=_SLATE_900)
            celda.alignment = Alignment(
                horizontal=(aligns[col - 1] if aligns else "left"), vertical="center", indent=1
            )
            if fmts and fmts[col - 1]:
                celda.number_format = fmts[col - 1]
        ws.row_dimensions[fila].height = 22
        return fila + 1

    def _fila_vacia(self, ws: Worksheet, fila: int, texto: str, ncols: int) -> int:
        ultima = get_column_letter(ncols)
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.font = Font(italic=True, size=11, color=_SLATE_500)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, ncols + 1):
            ws.cell(row=fila, column=col).border = _BORDE
        ws.row_dimensions[fila].height = 24
        return fila + 1

    def _nota(self, ws: Worksheet, fila: int, texto: str, ncols: int) -> int:
        ultima = get_column_letter(ncols)
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.font = Font(italic=True, size=9, color=_SLATE_500)
        celda.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        ws.row_dimensions[fila].height = 30
        return fila + 1

    def _anchos(self, ws: Worksheet, anchos: dict[str, float]) -> None:
        for letra, ancho in anchos.items():
            ws.column_dimensions[letra].width = ancho

    # -- Reportes -------------------------------------------------------------

    def _render_resumen(self, ws: Worksheet, datos: dict) -> None:
        ws.title = "Resumen de ventas"
        ncols = 3
        self._anchos(ws, {"A": 34, "B": 20, "C": 18})

        fila = self._cabecera(ws, datos.get("titulo", "Reporte de Ventas"), datos.get("periodo", ""), ncols)

        total_vendido = float(datos.get("total_vendido", 0) or 0)
        total_egresos = float(datos.get("total_egresos", 0) or 0)
        total_devuelto = float(datos.get("total_devuelto", 0) or 0)
        utilidad = float(datos.get("utilidad_neta", total_vendido - total_egresos))
        margen = (utilidad / total_vendido) if total_vendido > 0 else 0.0
        color_utilidad = _INDIGO if utilidad >= 0 else _AMBER

        fila = self._seccion(ws, fila, "Resumen del período", ncols)
        fila = self._kpi(ws, fila, "Total vendido (neto)", total_vendido, ncols, _MONEDA, _EMERALD, 0)
        fila = self._kpi(ws, fila, "Total egresos", total_egresos, ncols, _MONEDA, _ROSE, 1)
        fila = self._kpi(ws, fila, "Utilidad neta", utilidad, ncols, _MONEDA, color_utilidad, 2)
        fila = self._kpi(ws, fila, "Margen de utilidad", margen, ncols, _PORCENTAJE, color_utilidad, 3)
        fila = self._kpi(ws, fila, "Devoluciones (ya deducidas)", total_devuelto, ncols, _MONEDA, _SLATE_700, 4)
        fila = self._kpi(ws, fila, "N.° de ventas", int(datos.get("numero_ventas", 0) or 0), ncols, _ENTERO, _SLATE_900, 5)
        fila = self._kpi(ws, fila, "Ticket promedio", float(datos.get("ticket_promedio", 0) or 0), ncols, _MONEDA, _SLATE_900, 6)
        fila += 1

        metodos = datos.get("metodos_pago") or {}
        if metodos:
            fila = self._seccion(ws, fila, "Cobrado por método de pago", ncols)
            fila = self._encabezado_tabla(ws, fila, ["Método", "Monto", "% del total"], ["left", "right", "right"])
            total_metodos = sum(float(v or 0) for v in metodos.values())
            ordenados = sorted(metodos.items(), key=lambda kv: float(kv[1] or 0), reverse=True)
            for idx, (metodo, monto) in enumerate(ordenados):
                monto = float(monto or 0)
                pct = (monto / total_metodos) if total_metodos > 0 else 0.0
                fila = self._fila_tabla(
                    ws, fila, [metodo, monto, pct], [None, _MONEDA, _PORCENTAJE], ["left", "right", "right"], idx
                )
            fila = self._fila_total(
                ws, fila, ["Total cobrado", total_metodos, 1.0 if total_metodos > 0 else 0.0],
                [None, _MONEDA, _PORCENTAJE], ["left", "right", "right"],
            )
            fila += 1

        top = datos.get("top_productos") or []
        fila = self._seccion(ws, fila, "Productos más vendidos", ncols)
        fila = self._encabezado_tabla(ws, fila, ["Producto", "Unidades", "Total"], ["left", "right", "right"])
        if top:
            total_top = 0.0
            for idx, prod in enumerate(top):
                total_prod = float(prod.get("total", 0) or 0)
                total_top += total_prod
                fila = self._fila_tabla(
                    ws, fila,
                    [prod.get("nombre", ""), int(prod.get("cantidad", 0) or 0), total_prod],
                    [None, _ENTERO, _MONEDA], ["left", "right", "right"], idx,
                )
            self._fila_total(
                ws, fila, ["Total productos", "", total_top], [None, None, _MONEDA], ["left", "right", "right"]
            )
        else:
            self._fila_vacia(ws, fila, "No hubo ventas en el período.", ncols)

    def _render_egresos(self, ws: Worksheet, datos: dict) -> None:
        ws.title = "Egresos"
        ncols = 4
        self._anchos(ws, {"A": 14, "B": 46, "C": 12, "D": 18})

        fila = self._cabecera(ws, datos.get("titulo", "Reporte de Egresos"), datos.get("periodo", ""), ncols)

        egresos = datos.get("egresos") or []
        total = float(datos.get("total_egresos", sum(float(e.get("monto", 0) or 0) for e in egresos)))
        unidades = sum(int(e.get("unidades", 0) or 0) for e in egresos)
        registros = int(datos.get("numero_registros", len(egresos)))

        fila = self._seccion(ws, fila, "Resumen de egresos", ncols, _ROSE_HEADER)
        fila = self._kpi(ws, fila, "Total de egresos", total, ncols, _MONEDA, _ROSE, 0)
        fila = self._kpi(ws, fila, "Ingresos aprobados", registros, ncols, _ENTERO, _SLATE_900, 1)
        fila = self._kpi(ws, fila, "Unidades ingresadas", unidades, ncols, _ENTERO, _SLATE_900, 2)
        fila += 1

        fila = self._seccion(ws, fila, "Detalle de egresos", ncols, _ROSE_HEADER)
        fila = self._encabezado_tabla(
            ws, fila, ["Fecha", "Concepto", "Unidades", "Monto"], ["left", "left", "right", "right"], _ROSE_HEADER
        )
        ws.freeze_panes = f"A{fila}"  # mantiene visibles las cabeceras al hacer scroll
        if egresos:
            for idx, egreso in enumerate(egresos):
                fila = self._fila_tabla(
                    ws, fila,
                    [
                        egreso.get("fecha", ""),
                        egreso.get("concepto", ""),
                        int(egreso.get("unidades", 0) or 0),
                        float(egreso.get("monto", 0) or 0),
                    ],
                    [None, None, _ENTERO, _MONEDA], ["left", "left", "right", "right"], idx,
                )
            fila = self._fila_total(
                ws, fila, ["", "Total egresos", unidades, total],
                [None, None, _ENTERO, _MONEDA], ["left", "left", "right", "right"],
            )
        else:
            fila = self._fila_vacia(ws, fila, "No hay egresos registrados en este período.", ncols)

        fila += 1
        self._nota(
            ws, fila,
            "Cada egreso es el costo (cantidad × precio de compra unitario) de una solicitud de "
            "ingreso aprobada. Refleja el COSTO de la mercadería que entró al inventario, no la forma "
            "de pago al proveedor (contado o crédito).",
            ncols,
        )

    def _render_mas_vendidos(self, ws: Worksheet, datos: dict) -> None:
        ws.title = "Más vendidos"
        ncols = 4
        self._anchos(ws, {"A": 8, "B": 42, "C": 14, "D": 18})

        fila = self._cabecera(ws, datos.get("titulo", "Productos Más Vendidos"), datos.get("periodo", ""), ncols)

        productos = datos.get("productos") or []
        fila = self._seccion(ws, fila, "Ranking de productos", ncols)
        fila = self._encabezado_tabla(
            ws, fila, ["#", "Producto", "Cantidad", "Total"], ["center", "left", "right", "right"]
        )
        if productos:
            total = 0.0
            for idx, prod in enumerate(productos):
                total_prod = float(prod.get("total", 0) or 0)
                total += total_prod
                fila = self._fila_tabla(
                    ws, fila,
                    [idx + 1, prod.get("nombre", ""), int(prod.get("cantidad", 0) or 0), total_prod],
                    [_ENTERO, None, _ENTERO, _MONEDA], ["center", "left", "right", "right"], idx,
                )
            self._fila_total(
                ws, fila, ["", "Total", "", total], [None, None, None, _MONEDA], ["center", "left", "right", "right"]
            )
        else:
            self._fila_vacia(ws, fila, "No hubo ventas en el período.", ncols)
